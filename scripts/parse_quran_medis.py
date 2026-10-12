import re
import json
from docx import Document

SURAH_MAPPING = {
    "Al-Fatihah": 1, "Al-Baqarah": 2, "Ali 'Imran": 3, "An-Nisa'": 4, "Al-Ma'idah": 5,
    "Al-An'am": 6, "Al-A'raf": 7, "Al-Anfal": 8, "At-Taubah": 9, "Yunus": 10,
    "Hud": 11, "Yusuf": 12, "Ar-Ra'd": 13, "Ibrahim": 14, "Al-Hijr": 15,
    "An-Nahl": 16, "Al-Isra'": 17, "Al-Kahf": 18, "Maryam": 19, "Taha": 20,
    "Al-Anbiya'": 21, "Al-Hajj": 22, "Al-Mu'minun": 23, "An-Nur": 24, "Al-Furqan": 25,
    "Asy-Syu'ara'": 26, "An-Naml": 27, "Al-Qasas": 28, "Al-'Ankabut": 29, "Ar-Rum": 30,
    "Luqman": 31, "As-Sajdah": 32, "Al-Ahzab": 33, "Saba'": 34, "Fatir": 35,
    "Yasin": 36, "As-Saffat": 37, "Sad": 38, "Az-Zumar": 39, "Gafir": 40,
    "Fussilat": 41, "Asy-Syura": 42, "Az-Zukhruf": 43, "Ad-Dukhan": 44, "Al-Jasiyah": 45,
    "Al-Ahqaf": 46, "Muhammad": 47, "Al-Fath": 48, "Al-Hujurat": 49, "Qaf": 50,
    "Az-Zariyat": 51, "At-Tur": 52, "An-Najm": 53, "Al-Qamar": 54, "Ar-Rahman": 55,
    "Al-Waqi'ah": 56, "Al-Hadid": 57, "Al-Mujadalah": 58, "Al-Hasyr": 59, "Al-Mumtahanah": 60,
    "As-Saff": 61, "Al-Jumu'ah": 62, "Al-Munafiqun": 63, "At-Tagabun": 64, "At-Talaq": 65,
    "At-Tahrim": 66, "Al-Mulk": 67, "Al-Qalam": 68, "Al-Haqqah": 69, "Al-Ma'arij": 70,
    "Nuh": 71, "Al-Jinn": 72, "Al-Muzzammil": 73, "Al-Muddassir": 74, "Al-Qiyamah": 75,
    "Al-Insan": 76, "Al-Mursalat": 77, "An-Naba'": 78, "An-Nazi'at": 79, "'Abasa": 80,
    "At-Takwir": 81, "Al-Infitar": 82, "Al-Mutaffifin": 83, "Al-Insyiqaq": 84, "Al-Buruj": 85,
    "At-Tariq": 86, "Al-A'la": 87, "Al-Ghasyiyah": 88, "Al-Fajr": 89, "Al-Balad": 90,
    "Asy-Syams": 91, "Al-Lail": 92, "Ad-Duha": 93, "Asy-Syarh": 94, "At-Tin": 95,
    "Al-'Alaq": 96, "Al-Qadr": 97, "Al-Bayyinah": 98, "Az-Zalzalah": 99, "Al-'Adiyat": 100,
    "Al-Qari'ah": 101, "At-Takasur": 102, "Al-'Asr": 103, "Al-Humazah": 104, "Al-Fil": 105,
    "Quraisy": 106, "Al-Ma'un": 107, "Al-Kausar": 108, "Al-Kafirun": 109, "An-Nasr": 110,
    "Al-Lahab": 111, "Al-Ikhlas": 112, "Al-Falaq": 113, "An-Nas": 114
}

# Deteksi kategori dari paragraf / header tabel
KATEGORI_PATTERNS = [
    (r"I\.?\s*RUHANI", "I_RUHANI", "I. RUHANI", "Tema seputar kesehatan jiwa, hati, akhlak, dan spiritual"),
    (r"II\.?\s*RAGAWI", "II_RAGAWI", "II. RAGAWI", "Tema seputar kesehatan fisik, makanan, minuman, dan tubuh"),
    (r"III\.?\s*HIKMAH\s*IBADAH", "III_HIKMAH_IBADAH", "III. HIKMAH IBADAH", "Tema seputar hikmah ibadah, salat, zakat"),
    (r"IV\.?\s*ALA\s*NABI", "IV_ALA_NABI", "IV. ALA NABI", "Tema seputar sunah dan kebiasaan Rasulullah"),
]

TYPO_FIXES = {
    "Ai-Maidah": "Al-Ma'idah", "Al-Miidah": "Al-Ma'idah", "AlMaidah": "Al-Ma'idah",
    "Al-Arat": "Al-A'raf", "Al-AYaf": "Al-A'raf", "Al-Atat": "Al-A'raf",
    "Al-Arr'am": "Al-An'am", "Al-Anilm": "Al-An'am", "Ai-Anilm": "Al-An'am",
    "Al-Araf": "Al-A'raf", "Al-Kalf": "Al-Kahf", "Ai-Yaht": "Al-Kahf",
    "Al-Anbrya": "Al-Anbiya'", "AD-ahf": "Al-Kahf", "Al￥": "Al-Isra'",
    "Alinfitar": "Al-Infitar", "Al-Gasyiyah": "Al-Ghasyiyah", "Al-Asr": "Al-'Asr",
    "Abasa": "'Abasa", "N4107": "An-Nisa'", "pnisg": "An-Nisa'",
    "6Ns": "An-Nisa'", "f6t6s8": "An-Nisa'", "6-Nis8": "An-Nisa'",
}

def clean_surah_name(name):
    if not name: return None
    clean = name.strip()
    for typo, correct in TYPO_FIXES.items():
        if typo in clean:
            clean = clean.replace(typo, correct)
    return SURAH_MAPPING.get(clean)

def parse_rujukan(rujukan_str):
    if not rujukan_str: return None, None
    normalized = rujukan_str.replace('，', ',').replace('：', ':').replace('.', ':').strip()
    match = re.search(r"([A-Za-z\s''\-]+)[,，]?\s*(\d+)\s*[:：]\s*(\d+)", normalized)
    if not match: return None, None
    nama_surah = match.group(1).strip()
    ayat = int(match.group(3))
    return clean_surah_name(nama_surah), ayat

def detect_kategori(text):
    for pattern, key, nama, _ in KATEGORI_PATTERNS:
        if re.search(pattern, text, re.IGNORECASE):
            return key, nama
    return None, None

def parse_docx(file_path):
    doc = Document(file_path)
    
    # Struktur per kategori
    result = {
        "metadata": {
            "judul_buku": "Al-Qur'an dan Kesehatan",
            "total_kategori": 4,
            "sumber_api": "https://equran.id/api/v2"
        },
        "kategori": {
            key: {
                "id": key,
                "nama_lengkap": nama,
                "deskripsi": desk,
                "jumlah_tema": 0,
                "tema": []
            }
            for _, key, nama, desk in KATEGORI_PATTERNS
        }
    }
    
    current_kategori_key = None
    
    # Scan paragraf untuk deteksi kategori
    for para in doc.paragraphs:
        text = para.text.strip()
        key, nama = detect_kategori(text)
        if key:
            current_kategori_key = key
    
    # Parse tabel
    for table in doc.tables:
        # Cek header tabel untuk deteksi kategori
        if table.rows:
            header_text = " ".join([cell.text for cell in table.rows[0].cells])
            key, _ = detect_kategori(header_text)
            if key:
                current_kategori_key = key
        
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            
            # Skip header
            if any(h in " ".join(cells).upper() for h in ["NO.", "TEMA", "HAL.", "SURAH"]):
                continue
            
            no = tema = halaman = rujukan = None
            for cell in cells:
                if not cell: continue
                if re.match(r'^\d+$', cell):
                    if no is None:
                        no = int(cell)
                    elif halaman is None and int(cell) > 10:
                        halaman = int(cell)
                elif re.search(r'[A-Za-z].*[，,].*\d+[:：]\d+', cell):
                    rujukan = cell
                elif len(cell) > 3 and tema is None and not cell.isdigit():
                    tema = cell
            
            if tema and rujukan and current_kategori_key:
                surah_nomor, ayat_nomor = parse_rujukan(rujukan)
                if surah_nomor:
                    result["kategori"][current_kategori_key]["tema"].append({
                        "no": no,
                        "tema": tema,
                        "halaman": halaman,
                        "surah_nomor": surah_nomor,
                        "ayat_nomor": ayat_nomor
                    })
    
    # Hitung jumlah tema per kategori
    for key in result["kategori"]:
        result["kategori"][key]["jumlah_tema"] = len(result["kategori"][key]["tema"])
    
    total = sum(k["jumlah_tema"] for k in result["kategori"].values())
    result["metadata"]["total_tema"] = total
    
    return result

if __name__ == "__main__":
    data = parse_docx("quran medis.docx")
    with open("quranMedis.json", "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    
    print("✅ Hasil parsing per kategori:")
    for key, kat in data["kategori"].items():
        print(f"  • {kat['nama_lengkap']}: {kat['jumlah_tema']} tema")
    print(f"  TOTAL: {data['metadata']['total_tema']} tema")